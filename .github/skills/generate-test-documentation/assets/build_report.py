"""
build_report.py  —  Construye el reporte QA en Word en UN solo proceso.

Uso:
  uv run --with python-docx python build_report.py <plantilla_base.docx> <datos.json> <carpeta_salida> [dd/MM/yyyy]

Hace TODO:
  1. Copia la plantilla base (portada, encabezado/logo, indice, estilos).
  2. Portada: reemplaza el titulo (aunque este en text boxes) por "Plan de Pruebas" + titulo HU.
  3. Fechas: actualiza portada + encabezado(s) a la fecha actual (o la indicada).
  4. Cuerpo: agrega las secciones (Introduccion, Alcance, Resumen, Resultados, Evidencias,
     Conclusiones, Recomendaciones) con las imagenes de evidencia ajustadas a la pagina.
Imprime la ruta final del .docx generado.
"""

import sys, os, re, json, shutil, datetime
from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn

if len(sys.argv) < 4:
    print("ERROR: faltan argumentos"); sys.exit(1)

template   = sys.argv[1]
data_json  = sys.argv[2]
out_dir    = sys.argv[3]
fecha      = sys.argv[4] if len(sys.argv) > 4 else datetime.date.today().strftime('%d/%m/%Y')

with open(data_json, 'r', encoding='utf-8-sig') as f:
    data = json.load(f)

plan_id   = str(data.get('PlanId', ''))
hu        = data.get('Hu') or {}
hu_title  = (hu.get('Title') or '').strip()
summary   = data.get('Summary') or {}
suites    = data.get('Suites') or []

# ---------- ruta de salida ----------
def sanitize(s):
    s = re.sub(r'[^\w\- ]', '', s, flags=re.UNICODE).strip()
    return re.sub(r'\s+', '-', s)

os.makedirs(out_dir, exist_ok=True)
fname = f"Reporte_QA_Plan-{plan_id}_{sanitize(hu_title)}_{datetime.date.today().strftime('%Y-%m-%d')}.docx"
output = os.path.join(out_dir, fname)
shutil.copyfile(template, output)

doc = Document(output)

date_re_full = re.compile(r'^\d{1,2}/\d{1,2}/\d{4}$')
date_re      = re.compile(r'\d{1,2}/\d{1,2}/\d{4}')

def para_text(p):
    return ''.join((t.text or '') for t in p.iter(qn('w:t')))

def run_sz(r):
    rpr = r.find(qn('w:rPr'))
    if rpr is None: return None
    sz = rpr.find(qn('w:sz'))
    return sz.get(qn('w:val')) if sz is not None else None

def set_line(p, text):
    ts = list(p.iter(qn('w:t')))
    if not ts: return
    ts[0].text = text
    ts[0].set(qn('xml:space'), 'preserve')
    for t in ts[1:]:
        t.text = ''

# ---------- 1) TITULO (portada, incluye text boxes) ----------
# Paragrafos "hoja" (sin w:p anidados) con runs de fuente grande y letras (no fecha).
title_paras = []
for p in doc.element.body.iter(qn('w:p')):
    if p.find('.//' + qn('w:p')) is not None:
        continue  # es contenedor (envuelve un text box), no una linea de titulo
    txt = para_text(p).strip()
    if not txt:
        continue
    big = any(run_sz(r) in ('40','44','48','52','56','60','72') for r in p.iter(qn('w:r')))
    if big and re.search(r'[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]', txt) and not date_re_full.match(txt.replace(' ', '')):
        title_paras.append((p, txt))

distinct = []
for _, t in title_paras:
    if t not in distinct:
        distinct.append(t)
line1_old = distinct[0] if len(distinct) >= 1 else None
line2_old = distinct[1] if len(distinct) >= 2 else None
for p, t in title_paras:
    if t == line1_old:
        set_line(p, "Plan de Pruebas")
    elif t == line2_old:
        set_line(p, hu_title)

# ---------- 2) FECHAS (portada + encabezados/pies) ----------
# quitar placeholders para que la fecha se muestre como valor real
for sp in list(doc.element.iter(qn('w:showingPlcHdr'))):
    parent = sp.getparent()
    if parent is not None:
        parent.remove(sp)

def replace_dates(element, new_date):
    ts = list(element.iter(qn('w:t')))
    if not ts: return
    texts = [t.text or '' for t in ts]
    starts, charmap, buf, pos = [], [], [], 0
    for i, tx in enumerate(texts):
        starts.append(pos)
        for ch in tx:
            buf.append(ch); charmap.append(i); pos += 1
    s = ''.join(buf)
    matches = list(date_re.finditer(s))
    if not matches: return
    newtext = {i: texts[i] for i in range(len(texts))}
    for m in matches:
        a, b = m.start(), m.end() - 1
        fr, lr = charmap[a], charmap[b]
        prefix = texts[fr][:a - starts[fr]]
        suffix = texts[lr][(b - starts[lr]) + 1:]
        if fr == lr:
            newtext[fr] = prefix + new_date + suffix
        else:
            newtext[fr] = prefix + new_date
            for k in range(fr + 1, lr):
                newtext[k] = ''
            newtext[lr] = suffix
    for i, t in enumerate(ts):
        t.text = newtext[i]
        t.set(qn('xml:space'), 'preserve')

replace_dates(doc.element.body, fecha)
for section in doc.sections:
    for hf in (section.header, section.footer,
               section.first_page_header, section.first_page_footer,
               section.even_page_header, section.even_page_footer):
        try:
            replace_dates(hf._element, fecha)
        except Exception:
            pass

# ---------- 3) CUERPO ----------
def heading1(text):
    p = doc.add_paragraph(text)
    pPr = p._p.get_or_add_pPr()
    st = pPr.makeelement(qn('w:pStyle'), {qn('w:val'): 'Ttulo1'})
    pPr.insert(0, st)
    return p

def bullet(text):
    p = doc.add_paragraph(text)
    pPr = p._p.get_or_add_pPr()
    st = pPr.makeelement(qn('w:pStyle'), {qn('w:val'): 'Prrafodelista'})
    pPr.insert(0, st)
    return p

def normal(text):
    return doc.add_paragraph(text)

def bold(text):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True
    return p

mark = {'passed': '\u2714\ufe0f', 'failed': '\u2717'}

# INTRODUCCION
heading1("INTRODUCCIÓN")
normal(f"Este documento presenta la validación funcional del plan de pruebas "
       f"\u201c{data.get('PlanName','')}\u201d, asociado a la Historia de Usuario #{hu.get('Id','')}.")
if hu.get('Description'):
    normal(hu['Description'])

# ALCANCE
heading1("ALCANCE DE PRUEBAS")
normal("Se probaron los siguientes casos, organizados por suite:")
for s in suites:
    bullet(s.get('SuiteName', ''))

# RESUMEN EJECUTIVO
heading1("RESUMEN EJECUTIVO")
normal("Resumen de la ejecución de los casos de prueba:")
bullet(f"Total de casos: {summary.get('total', 0)}")
bullet(f"Casos exitosos: {summary.get('passed', 0)} ({summary.get('passRate', 0)}%)")
bullet(f"Casos fallidos: {summary.get('failed', 0)}")
bullet(f"Casos bloqueados: {summary.get('blocked', 0)}")
bullet(f"Casos no aplicables: {summary.get('notApplicable', 0)}")
bullet(f"Casos sin ejecutar: {summary.get('unspecified', 0)}")

# RESULTADOS DE PRUEBAS
heading1("RESULTADOS DE PRUEBAS")
for s in suites:
    bold(s.get('SuiteName', ''))
    for c in s.get('Cases', []):
        oc = str(c.get('Outcome', '')).lower()
        m = mark.get(oc, f"({c.get('Outcome','')})")
        bullet(f"{c.get('TcName','')} {m}")

# EVIDENCIAS
heading1("EVIDENCIAS")
any_evi = False
for s in suites:
    for c in s.get('Cases', []):
        paths = c.get('EvidencePaths') or []
        if not paths:
            continue
        any_evi = True
        bold(c.get('TcName', ''))
        if c.get('Comment'):
            normal(f"Comentario de ejecución: {c['Comment']}")
        for ph in paths:
            if os.path.exists(ph):
                try:
                    doc.add_picture(ph, width=Inches(6))
                except Exception:
                    normal(f"[No se pudo insertar la imagen: {os.path.basename(ph)}]")
if not any_evi:
    normal("No se registraron evidencias para esta ejecución.")

# CONCLUSIONES
heading1("CONCLUSIONES")
if summary.get('failed', 0) == 0 and summary.get('passed', 0) > 0:
    normal("La ejecución del plan de pruebas se completó satisfactoriamente. "
           "Los casos ejecutados cumplieron con el comportamiento esperado.")
else:
    normal(f"Durante la ejecución del plan se registraron {summary.get('failed',0)} caso(s) fallido(s) "
           f"de un total de {summary.get('total',0)}. Se recomienda revisar los defectos reportados "
           f"antes de dar por finalizada la validación.")

# RECOMENDACIONES
heading1("RECOMENDACIONES")
if summary.get('failed', 0) > 0:
    normal("Atender los casos fallidos y sus defectos asociados, y volver a ejecutar las pruebas "
           "correspondientes para verificar su corrección.")
if summary.get('unspecified', 0) > 0:
    normal(f"Ejecutar los {summary.get('unspecified',0)} caso(s) que quedaron sin ejecutar para "
           f"completar la cobertura del plan.")
if summary.get('failed', 0) == 0 and summary.get('unspecified', 0) == 0:
    normal("Mantener la cobertura de pruebas en futuras iteraciones del desarrollo.")

doc.save(output)
print(output)
