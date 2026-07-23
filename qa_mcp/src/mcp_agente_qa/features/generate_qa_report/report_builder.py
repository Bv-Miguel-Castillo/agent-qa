"""Builds the QA Word report in-process (ported from the old build_report.py script).

Runs entirely inside the server: receives the collected plan data (dict) and returns the
path of the generated .docx. No local/agent command execution is involved.
"""

from __future__ import annotations

import datetime
import os
import re
import shutil

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches


def _sanitize(text: str) -> str:
    text = re.sub(r"[^\w\- ]", "", text, flags=re.UNICODE).strip()
    return re.sub(r"\s+", "-", text)


def build_report_docx(
    data: dict,
    template_path: str,
    out_dir: str,
    fecha: str | None = None,
) -> str:
    """Build the QA Word report from the collected plan data and return the .docx path."""
    fecha = fecha or datetime.date.today().strftime("%d/%m/%Y")

    plan_id = str(data.get("PlanId", ""))
    hu = data.get("Hu") or {}
    hu_title = (hu.get("Title") or "").strip()
    summary = data.get("Summary") or {}
    suites = data.get("Suites") or []

    os.makedirs(out_dir, exist_ok=True)
    fname = (
        f"Reporte_QA_Plan-{plan_id}_{_sanitize(hu_title)}_"
        f"{datetime.date.today().strftime('%Y-%m-%d')}.docx"
    )
    output = os.path.join(out_dir, fname)
    shutil.copyfile(template_path, output)

    doc = Document(output)

    date_re_full = re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$")
    date_re = re.compile(r"\d{1,2}/\d{1,2}/\d{4}")

    def para_text(p):
        return "".join((t.text or "") for t in p.iter(qn("w:t")))

    def run_sz(r):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            return None
        sz = rpr.find(qn("w:sz"))
        return sz.get(qn("w:val")) if sz is not None else None

    def set_line(p, text):
        ts = list(p.iter(qn("w:t")))
        if not ts:
            return
        ts[0].text = text
        ts[0].set(qn("xml:space"), "preserve")
        for t in ts[1:]:
            t.text = ""

    # 1) Title (cover, includes text boxes)
    title_paras = []
    for p in doc.element.body.iter(qn("w:p")):
        if p.find(".//" + qn("w:p")) is not None:
            continue
        txt = para_text(p).strip()
        if not txt:
            continue
        big = any(
            run_sz(r) in ("40", "44", "48", "52", "56", "60", "72") for r in p.iter(qn("w:r"))
        )
        if big and re.search(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ]", txt) and not date_re_full.match(
            txt.replace(" ", "")
        ):
            title_paras.append((p, txt))

    distinct = []
    for _, t in title_paras:
        if t not in distinct:
            distinct.append(t)
    line1_old = distinct[0] if len(distinct) >= 1 else None
    line2_old = distinct[1] if len(distinct) >= 2 else None
    for p, t in title_paras:
        if t == line1_old:
            set_line(p, "Plan de Pruebas")
        elif t == line2_old:
            set_line(p, hu_title)

    # 2) Dates (cover + headers/footers)
    for sp in list(doc.element.iter(qn("w:showingPlcHdr"))):
        parent = sp.getparent()
        if parent is not None:
            parent.remove(sp)

    def replace_dates(element, new_date):
        ts = list(element.iter(qn("w:t")))
        if not ts:
            return
        texts = [t.text or "" for t in ts]
        starts, charmap, buf, pos = [], [], [], 0
        for i, tx in enumerate(texts):
            starts.append(pos)
            for ch in tx:
                buf.append(ch)
                charmap.append(i)
                pos += 1
        s = "".join(buf)
        matches = list(date_re.finditer(s))
        if not matches:
            return
        newtext = {i: texts[i] for i in range(len(texts))}
        for m in matches:
            a, b = m.start(), m.end() - 1
            fr, lr = charmap[a], charmap[b]
            prefix = texts[fr][: a - starts[fr]]
            suffix = texts[lr][(b - starts[lr]) + 1 :]
            if fr == lr:
                newtext[fr] = prefix + new_date + suffix
            else:
                newtext[fr] = prefix + new_date
                for k in range(fr + 1, lr):
                    newtext[k] = ""
                newtext[lr] = suffix
        for i, t in enumerate(ts):
            t.text = newtext[i]
            t.set(qn("xml:space"), "preserve")

    replace_dates(doc.element.body, fecha)
    for section in doc.sections:
        for hf in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            try:
                replace_dates(hf._element, fecha)
            except Exception:
                pass

    # 3) Body
    def heading1(text):
        p = doc.add_paragraph(text)
        pPr = p._p.get_or_add_pPr()
        st = pPr.makeelement(qn("w:pStyle"), {qn("w:val"): "Ttulo1"})
        pPr.insert(0, st)
        return p

    def bullet(text):
        p = doc.add_paragraph(text)
        pPr = p._p.get_or_add_pPr()
        st = pPr.makeelement(qn("w:pStyle"), {qn("w:val"): "Prrafodelista"})
        pPr.insert(0, st)
        return p

    def normal(text):
        return doc.add_paragraph(text)

    def bold(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.bold = True
        return p

    mark = {"passed": "\u2714\ufe0f", "failed": "\u2717"}

    heading1("INTRODUCCIÓN")
    normal(
        f"Este documento presenta la validación funcional del plan de pruebas "
        f"\u201c{data.get('PlanName', '')}\u201d, asociado a la Historia de Usuario "
        f"#{hu.get('Id', '')}."
    )
    if hu.get("Description"):
        normal(hu["Description"])

    heading1("ALCANCE DE PRUEBAS")
    normal("Se probaron los siguientes casos, organizados por suite:")
    for s in suites:
        bullet(s.get("SuiteName", ""))

    heading1("RESUMEN EJECUTIVO")
    normal("Resumen de la ejecución de los casos de prueba:")
    bullet(f"Total de casos: {summary.get('total', 0)}")
    bullet(f"Casos exitosos: {summary.get('passed', 0)} ({summary.get('passRate', 0)}%)")
    bullet(f"Casos fallidos: {summary.get('failed', 0)}")
    bullet(f"Casos bloqueados: {summary.get('blocked', 0)}")
    bullet(f"Casos no aplicables: {summary.get('notApplicable', 0)}")
    bullet(f"Casos sin ejecutar: {summary.get('unspecified', 0)}")

    heading1("RESULTADOS DE PRUEBAS")
    for s in suites:
        bold(s.get("SuiteName", ""))
        for c in s.get("Cases", []):
            oc = str(c.get("Outcome", "")).lower()
            m = mark.get(oc, f"({c.get('Outcome', '')})")
            bullet(f"{c.get('TcName', '')} {m}")

    heading1("EVIDENCIAS")
    any_item = False
    for s in suites:
        for c in s.get("Cases", []):
            paths = c.get("EvidencePaths") or []
            comment = c.get("Comment")
            # Include the case if it has a comment OR evidence images (so comments are never lost).
            if not paths and not comment:
                continue
            any_item = True
            bold(c.get("TcName", ""))
            if comment:
                normal(f"Comentario de ejecución: {comment}")
            if paths:
                for ph in paths:
                    if os.path.exists(ph):
                        try:
                            doc.add_picture(ph, width=Inches(6))
                        except Exception:
                            normal(f"[No se pudo insertar la imagen: {os.path.basename(ph)}]")
            else:
                normal("No se encontraron imágenes de evidencia para este caso.")
    if not any_item:
        normal("No se registraron evidencias para esta ejecución.")

    heading1("CONCLUSIONES")
    if summary.get("failed", 0) == 0 and summary.get("passed", 0) > 0:
        normal(
            "La ejecución del plan de pruebas se completó satisfactoriamente. "
            "Los casos ejecutados cumplieron con el comportamiento esperado."
        )
    else:
        normal(
            f"Durante la ejecución del plan se registraron {summary.get('failed', 0)} caso(s) "
            f"fallido(s) de un total de {summary.get('total', 0)}. Se recomienda revisar los "
            f"defectos reportados antes de dar por finalizada la validación."
        )

    heading1("RECOMENDACIONES")
    if summary.get("failed", 0) > 0:
        normal(
            "Atender los casos fallidos y sus defectos asociados, y volver a ejecutar las pruebas "
            "correspondientes para verificar su corrección."
        )
    if summary.get("unspecified", 0) > 0:
        normal(
            f"Ejecutar los {summary.get('unspecified', 0)} caso(s) que quedaron sin ejecutar para "
            f"completar la cobertura del plan."
        )
    if summary.get("failed", 0) == 0 and summary.get("unspecified", 0) == 0:
        normal("Mantener la cobertura de pruebas en futuras iteraciones del desarrollo.")

    doc.save(output)
    return output
